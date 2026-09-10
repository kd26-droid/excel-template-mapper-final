import gzip
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from .mpn_pattern_library import (
    assess_mpn_text,
    character_classes,
    exact_mpn_pattern,
    grouped_mpn_pattern,
)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip()


def norm_words(value):
    return re.sub(r"[^A-Z0-9]+", " ", clean(value).upper()).strip()


def norm_compact(value):
    return re.sub(r"[^A-Z0-9]+", "", clean(value).upper())


def read_json(path, default):
    if not path.exists():
        return default
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_json(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)


def read_json_gz(path, default):
    if not path.exists():
        return default
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        return json.load(handle)


def write_json_gz(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(path, "wt", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))


def append_example(examples, key, value, max_examples=5):
    if not key or not value:
        return
    bucket = examples[key]
    if value not in bucket and len(bucket) < max_examples:
        bucket.append(value)


def find_column(headers, patterns, fallback_index):
    for index, header in enumerate(headers):
        text = norm_words(header)
        if any(pattern.search(text) for pattern in patterns):
            return index
    return fallback_index


def import_mpn_mfr_master(input_file, output_dir, sheet_name="MPN-MFR"):
    input_path = Path(input_file).expanduser()
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    output_path = Path(output_dir).expanduser()
    output_path.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Worksheet not found: {sheet_name}")

    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows, [])]
    if not headers:
        raise ValueError(f"Worksheet is empty: {sheet_name}")

    mpn_index = find_column(
        headers,
        [re.compile(r"MANUFACTURER PART"), re.compile(r"\bMPN\b"), re.compile(r"PART NU")],
        0,
    )
    mfr_index = find_column(
        headers,
        [re.compile(r"MANUFACTURER NAME"), re.compile(r"\bMFR\b"), re.compile(r"\bMFG\b")],
        1,
    )

    raw_rows = 0
    skipped_generic_specs = 0
    unique_mpns = {}
    pair_counts = Counter()
    manufacturer_counts = Counter()

    for row in rows:
        cells = list(row or [])
        mpn = clean(cells[mpn_index] if mpn_index < len(cells) else "")
        manufacturer = clean(cells[mfr_index] if mfr_index < len(cells) else "")
        if not mpn and not manufacturer:
            continue

        raw_rows += 1
        mpn_assessment = assess_mpn_text(mpn)
        if mpn_assessment["hard_reject"]:
            skipped_generic_specs += 1
            continue

        mpn = mpn_assessment["candidate_text"]
        mpn_key = norm_compact(mpn)
        manufacturer_key = norm_words(manufacturer)
        if not mpn_key:
            continue

        unique_mpns.setdefault(mpn_key, mpn)
        if manufacturer_key:
            pair_counts[(mpn_key, manufacturer)] += 1
            manufacturer_counts[manufacturer] += 1

    manufacturers_path = output_path / "manufacturers.json"
    manufacturers = read_json(manufacturers_path, {
        "version": 1,
        "source": "",
        "entry_count": 0,
        "name_count": 0,
        "alias_count": 0,
        "entries": [],
        "names": [],
        "aliases": {},
    })
    entries = list(manufacturers.get("entries") or [])
    names = list(manufacturers.get("names") or [])
    aliases = dict(manufacturers.get("aliases") or {})
    known_manufacturer_keys = {norm_compact(name) for name in names}
    known_manufacturer_keys.update(norm_compact(key) for key in aliases.keys())

    added_manufacturers = 0
    for manufacturer in sorted(manufacturer_counts, key=norm_words):
        key = norm_compact(manufacturer)
        if not key or key in known_manufacturer_keys:
            continue
        entries.append({"name": manufacturer, "aliases": [manufacturer]})
        names.append(manufacturer)
        known_manufacturer_keys.add(key)
        added_manufacturers += 1

    manufacturers.update({
        "version": 1,
        "source": ", ".join(filter(None, [str(manufacturers.get("source") or ""), input_path.name])),
        "entry_count": len(entries),
        "name_count": len(names),
        "alias_count": len(aliases),
        "entries": entries,
        "names": names,
        "aliases": aliases,
    })
    write_json(manufacturers_path, manufacturers)

    lookup_path = output_path / "mpn_mfr_lookup.v1.json.gz"
    previous_lookup = read_json_gz(lookup_path, {
        "version": 1,
        "sources": [],
        "row_count": 0,
        "unique_mpn_count": 0,
        "pair_count": 0,
        "manufacturers_by_mpn": {},
    })
    loaded_previous_by_mpn = previous_lookup.get("manufacturers_by_mpn") or {}
    previous_by_mpn = {}
    for key, value in loaded_previous_by_mpn.items():
        raw_mpn = value.get("mpn", key) if isinstance(value, dict) else key
        mpn_assessment = assess_mpn_text(raw_mpn)
        if mpn_assessment["hard_reject"]:
            continue

        canonical_mpn = mpn_assessment["candidate_text"]
        canonical_key = norm_compact(canonical_mpn)
        if not canonical_key:
            continue

        incoming = dict(value) if isinstance(value, dict) else {"mpn": canonical_mpn}
        incoming["mpn"] = canonical_mpn
        incoming.setdefault("manufacturers", [])
        incoming.setdefault("counts", {})

        existing = previous_by_mpn.setdefault(canonical_key, {
            "mpn": canonical_mpn,
            "manufacturers": [],
            "counts": {},
        })
        for manufacturer in incoming.get("manufacturers") or []:
            if manufacturer not in existing["manufacturers"]:
                existing["manufacturers"].append(manufacturer)
        for manufacturer, count in (incoming.get("counts") or {}).items():
            existing["counts"][manufacturer] = int(existing["counts"].get(manufacturer, 0) or 0) + int(count or 0)
    removed_existing_generic_specs = len(loaded_previous_by_mpn) - len(previous_by_mpn)
    manufacturers_by_mpn = dict(previous_by_mpn)
    lookup_sources = list(previous_lookup.get("sources") or [])
    source_already_imported = input_path.name in lookup_sources
    if not source_already_imported:
        lookup_sources.append(input_path.name)

    new_mpn_keys = set()
    new_pairs = 0
    for mpn_key, mpn in unique_mpns.items():
        if mpn_key not in previous_by_mpn:
            new_mpn_keys.add(mpn_key)
        manufacturers_by_mpn.setdefault(mpn_key, {
            "mpn": mpn,
            "manufacturers": [],
            "counts": {},
        })

    for (mpn_key, manufacturer), count in pair_counts.items():
        entry = manufacturers_by_mpn.setdefault(mpn_key, {
            "mpn": unique_mpns.get(mpn_key, mpn_key),
            "manufacturers": [],
            "counts": {},
        })
        entry.setdefault("mpn", unique_mpns.get(mpn_key, mpn_key))
        entry.setdefault("manufacturers", [])
        entry.setdefault("counts", {})
        if manufacturer not in entry["manufacturers"]:
            entry["manufacturers"].append(manufacturer)
            new_pairs += 1
        entry["counts"][manufacturer] = int(entry["counts"].get(manufacturer, 0) or 0) + int(count)

    for entry in manufacturers_by_mpn.values():
        entry["manufacturers"] = sorted(entry.get("manufacturers") or [], key=str.casefold)

    lookup_payload = {
        "version": 1,
        "sources": lookup_sources,
        "row_count": (
            max(int(previous_lookup.get("row_count", 0) or 0), raw_rows)
            if source_already_imported
            else int(previous_lookup.get("row_count", 0) or 0) + raw_rows
        ),
        "unique_mpn_count": len(manufacturers_by_mpn),
        "pair_count": sum(len(entry.get("manufacturers") or []) for entry in manufacturers_by_mpn.values()),
        "manufacturers_by_mpn": manufacturers_by_mpn,
    }
    write_json_gz(lookup_path, lookup_payload)

    pattern_path = output_path / "mpn_patterns.v1.json.gz"
    pattern_payload = read_json_gz(pattern_path, {
        "version": 1,
        "source": "",
        "unique_mpn_rows": 0,
        "pattern_rule": "Letters become A, digits become 0, and spaces/symbols are retained. Grouped patterns compress A/0 runs.",
        "exact_patterns": {},
        "grouped_patterns": {},
        "character_classes": {},
        "lengths": {},
        "manufacturer_coverage": {},
        "grouped_examples": {},
        "character_class_examples": {},
    })
    exact_counts = Counter(pattern_payload.get("exact_patterns") or {})
    grouped_counts = Counter(pattern_payload.get("grouped_patterns") or {})
    class_counts = Counter(pattern_payload.get("character_classes") or {})
    length_counts = Counter(pattern_payload.get("lengths") or {})
    grouped_examples = defaultdict(list, {
        key: list(value)
        for key, value in (pattern_payload.get("grouped_examples") or {}).items()
    })
    class_examples = defaultdict(list, {
        key: list(value)
        for key, value in (pattern_payload.get("character_class_examples") or {}).items()
    })

    for mpn_key in sorted(new_mpn_keys):
        mpn = unique_mpns.get(mpn_key, mpn_key)
        exact = exact_mpn_pattern(mpn)
        grouped = grouped_mpn_pattern(mpn)
        cls = character_classes(mpn)
        exact_counts[exact] += 1
        grouped_counts[grouped] += 1
        class_counts[cls] += 1
        length_counts[str(len(mpn))] += 1
        append_example(grouped_examples, grouped, mpn)
        append_example(class_examples, cls, mpn, max_examples=3)

    pattern_source = str(pattern_payload.get("source") or "")
    if input_path.name not in pattern_source:
        pattern_source = ", ".join(filter(None, [pattern_source, input_path.name]))
    pattern_payload.update({
        "version": 1,
        "source": pattern_source,
        "unique_mpn_rows": int(pattern_payload.get("unique_mpn_rows", 0) or 0) + len(new_mpn_keys),
        "exact_patterns": dict(exact_counts),
        "grouped_patterns": dict(grouped_counts),
        "character_classes": dict(class_counts),
        "lengths": dict(length_counts),
        "grouped_examples": dict(grouped_examples),
        "character_class_examples": dict(class_examples),
    })
    write_json_gz(pattern_path, pattern_payload)

    manufacturer_pattern_path = output_path / "manufacturer_pattern_index.v1.json.gz"
    manufacturer_payload = read_json_gz(manufacturer_pattern_path, {
        "version": 1,
        "source": "",
        "manufacturer_rows": 0,
        "manufacturer_count": 0,
        "manufacturers": {},
    })
    manufacturer_index = dict(manufacturer_payload.get("manufacturers") or {})
    manufacturer_counters = defaultdict(Counter)
    manufacturer_examples = defaultdict(lambda: defaultdict(list))

    for manufacturer, data in manufacturer_index.items():
        for item in data.get("top_grouped_patterns") or []:
            pattern = item.get("pattern")
            if not pattern:
                continue
            manufacturer_counters[manufacturer][pattern] += int(item.get("count", 0) or 0)
            manufacturer_examples[manufacturer][pattern].extend(item.get("examples") or [])

    for (mpn_key, manufacturer), count in pair_counts.items():
        if mpn_key not in new_mpn_keys:
            continue
        mpn = unique_mpns.get(mpn_key, mpn_key)
        grouped = grouped_mpn_pattern(mpn)
        manufacturer_counters[manufacturer][grouped] += int(count)
        append_example(manufacturer_examples[manufacturer], grouped, mpn, max_examples=3)

    next_manufacturer_index = {}
    for manufacturer, counter in manufacturer_counters.items():
        top_patterns = []
        for pattern, count in counter.most_common(100):
            examples = []
            for example in manufacturer_examples[manufacturer].get(pattern, []):
                if example not in examples:
                    examples.append(example)
                if len(examples) >= 3:
                    break
            top_patterns.append({
                "pattern": pattern,
                "count": count,
                "examples": examples,
            })
        next_manufacturer_index[manufacturer] = {
            "mpn_count": sum(counter.values()),
            "top_grouped_patterns": top_patterns,
        }

    manufacturer_pattern_source = str(manufacturer_payload.get("source") or "")
    if input_path.name not in manufacturer_pattern_source:
        manufacturer_pattern_source = ", ".join(filter(None, [manufacturer_pattern_source, input_path.name]))
    manufacturer_pattern_payload = {
        "version": 1,
        "source": manufacturer_pattern_source,
        "manufacturer_rows": int(manufacturer_payload.get("manufacturer_rows", 0) or 0) + (0 if source_already_imported else sum(pair_counts.values())),
        "manufacturer_count": len(next_manufacturer_index),
        "manufacturers": next_manufacturer_index,
    }
    write_json_gz(manufacturer_pattern_path, manufacturer_pattern_payload)

    return {
        "raw_rows": raw_rows,
        "skipped_generic_specs": skipped_generic_specs,
        "removed_existing_generic_specs": removed_existing_generic_specs,
        "unique_mpns": len(unique_mpns),
        "unique_manufacturers": len(manufacturer_counts),
        "added_manufacturers": added_manufacturers,
        "new_exact_mpns": len(new_mpn_keys),
        "new_pairs": new_pairs,
        "lookup_path": str(lookup_path),
        "pattern_path": str(pattern_path),
        "manufacturer_pattern_path": str(manufacturer_pattern_path),
        "manufacturers_path": str(manufacturers_path),
    }
