import json
import re
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError


GENERIC_ALIAS_KEYS = {
    "and",
    "any",
    "component",
    "components",
    "corp",
    "corporation",
    "electronic",
    "electronics",
    "inc",
    "limited",
    "ltd",
    "mfr",
    "semiconductor",
    "technology",
}


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip()


def norm_words(value):
    return re.sub(r"[^a-z0-9]+", " ", clean(value).lower()).strip()


def norm_compact(value):
    return re.sub(r"[^a-z0-9]+", "", clean(value).lower())


def unique(values):
    seen = set()
    output = []
    for value in values:
        text = clean(value)
        key = norm_compact(text)
        if not text or not key or key in seen:
            continue
        seen.add(key)
        output.append(text)
    return output


def looks_like_alias(candidate, target_values):
    candidate_text = clean(candidate)
    candidate_key = norm_compact(candidate_text)
    candidate_words = norm_words(candidate_text).split()
    if not candidate_key or candidate_key in GENERIC_ALIAS_KEYS:
        return False

    for target in target_values:
        target_text = clean(target)
        target_key = norm_compact(target_text)
        target_words = norm_words(target_text).split()
        if not target_key:
            continue
        if candidate_key == target_key:
            return True
        if len(candidate_key) >= 3 and target_key.startswith(candidate_key):
            return True
        if len(candidate_key) >= 3 and len(target_key) >= 5 and candidate_key.startswith(target_key):
            return True
        if target_words and len(candidate_key) <= 8 and target_words[0].startswith(candidate_key):
            return True
        if candidate_words and target_words and len(candidate_words) <= min(3, len(target_words)):
            if all(
                target_words[index].startswith(word)
                for index, word in enumerate(candidate_words)
            ):
                return True
    return False


def workbook_rows(path, sheet_name=None):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    for row in worksheet.iter_rows(values_only=True):
        values = unique(row)
        if values:
            yield values


def parse_manufacturer_rows(rows):
    entries = []
    pending_single = None
    merged_singletons = 0
    index = 0

    while index < len(rows):
        values = rows[index]

        if len(values) == 1:
            if pending_single:
                entries.append({
                    "name": pending_single,
                    "aliases": [pending_single],
                })
            pending_single = values[0]
            index += 1
            continue

        aliases = list(values)
        if pending_single:
            if looks_like_alias(pending_single, aliases):
                aliases.insert(1, pending_single)
                merged_singletons += 1
            else:
                entries.append({
                    "name": pending_single,
                    "aliases": [pending_single],
                })
            pending_single = None

        lookahead = index + 1
        while lookahead < len(rows) and len(rows[lookahead]) == 1 and looks_like_alias(rows[lookahead][0], aliases):
            aliases.append(rows[lookahead][0])
            merged_singletons += 1
            lookahead += 1

        aliases = unique(aliases)
        entries.append({
            "name": aliases[0],
            "aliases": aliases,
        })
        index = lookahead

    if pending_single:
        entries.append({
            "name": pending_single,
            "aliases": [pending_single],
        })

    return entries, merged_singletons


def build_directory(entries, source):
    names = []
    aliases = {}
    seen_names = set()

    for entry in entries:
        canonical = clean(entry["name"])
        if not canonical:
            continue
        canonical_key = norm_compact(canonical)
        if canonical_key not in seen_names:
            names.append(canonical)
            seen_names.add(canonical_key)
        for alias in entry.get("aliases", []):
            alias_text = clean(alias)
            alias_key = norm_words(alias_text).upper()
            if alias_text and alias_key and alias_text != canonical:
                aliases[alias_key] = canonical

    return {
        "version": 1,
        "source": str(source),
        "entry_count": len(entries),
        "name_count": len(names),
        "alias_count": len(aliases),
        "entries": entries,
        "names": names,
        "aliases": aliases,
    }


class Command(BaseCommand):
    help = "Import manufacturer master Excel/CSV-style rows into backend JSON."

    def add_arguments(self, parser):
        parser.add_argument("input_file", help="Path to manufacturer .xlsx file")
        parser.add_argument("--sheet", default=None, help="Worksheet name. Defaults to first sheet.")
        parser.add_argument(
            "--output",
            default=None,
            help="Output JSON path. Defaults to excel_mapper/data/manufacturers.json",
        )

    def handle(self, *args, **options):
        input_path = Path(options["input_file"]).expanduser()
        if not input_path.exists():
            raise CommandError(f"Input file not found: {input_path}")

        output_path = Path(options["output"]).expanduser() if options["output"] else (
            Path(__file__).resolve().parents[2] / "data" / "manufacturers.json"
        )
        output_path.parent.mkdir(parents=True, exist_ok=True)

        rows = list(workbook_rows(input_path, options.get("sheet")))
        entries, merged_singletons = parse_manufacturer_rows(rows)
        directory = build_directory(entries, input_path.name)

        with output_path.open("w", encoding="utf-8") as handle:
            json.dump(directory, handle, ensure_ascii=False, indent=2)

        self.stdout.write(self.style.SUCCESS(
            f"Imported {directory['entry_count']} manufacturer entries, "
            f"{directory['alias_count']} aliases, merged {merged_singletons} adjacent singleton aliases."
        ))
        self.stdout.write(str(output_path))
