import gzip
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError

from excel_mapper.services.mpn_pattern_library import (
    character_classes,
    exact_mpn_pattern,
    grouped_mpn_pattern,
)


def clean(value):
    return re.sub(r"\s+", " ", str(value or "").replace("\u00a0", " ")).strip()


def write_json_gz(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(path, "wt", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))


def add_example(examples, key, value, max_examples=5):
    value = clean(value)
    if not key or not value:
        return
    bucket = examples[key]
    if value not in bucket and len(bucket) < max_examples:
        bucket.append(value)


class Command(BaseCommand):
    help = "Build compact MPN pattern-library JSON assets from the MPN/manufacturer pattern workbook."

    def add_arguments(self, parser):
        parser.add_argument("input_file", help="Path to Unique_MPN_and_Manufacturer_Patterns.xlsx")
        parser.add_argument(
            "--output-dir",
            default=None,
            help="Output directory. Defaults to backend/excel_mapper/data",
        )

    def handle(self, *args, **options):
        input_path = Path(options["input_file"]).expanduser()
        if not input_path.exists():
            raise CommandError(f"Input file not found: {input_path}")

        output_dir = Path(options["output_dir"]).expanduser() if options["output_dir"] else (
            Path(__file__).resolve().parents[2] / "data"
        )
        output_dir.mkdir(parents=True, exist_ok=True)

        workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
        required_sheets = {"Unique MPN Patterns", "Manufacturer MPN Patterns"}
        missing = sorted(required_sheets - set(workbook.sheetnames))
        if missing:
            raise CommandError(f"Workbook missing required sheet(s): {', '.join(missing)}")

        exact_counts = Counter()
        grouped_counts = Counter()
        class_counts = Counter()
        length_counts = Counter()
        manufacturer_coverage = Counter()
        grouped_examples = defaultdict(list)
        class_examples = defaultdict(list)
        total_unique_rows = 0

        unique_ws = workbook["Unique MPN Patterns"]
        for row in unique_ws.iter_rows(min_row=2, values_only=True):
            mpn, exact, grouped, length, cls, manufacturer_count, manufacturers = (list(row) + [None] * 7)[:7]
            mpn = clean(mpn)
            if not mpn:
                continue
            exact = clean(exact) or exact_mpn_pattern(mpn)
            grouped = clean(grouped) or grouped_mpn_pattern(mpn)
            cls = clean(cls) or character_classes(mpn)
            try:
                length_value = int(float(length))
            except (TypeError, ValueError):
                length_value = len(mpn)
            try:
                manufacturer_count_value = int(float(manufacturer_count))
            except (TypeError, ValueError):
                manufacturer_count_value = 0

            total_unique_rows += 1
            exact_counts[exact] += 1
            grouped_counts[grouped] += 1
            class_counts[cls] += 1
            length_counts[str(length_value)] += 1
            manufacturer_coverage[str(manufacturer_count_value)] += 1
            add_example(grouped_examples, grouped, mpn)
            add_example(class_examples, cls, mpn, max_examples=3)

        manufacturer_rows = 0
        manufacturer_index = {}
        manufacturer_grouped_counts = defaultdict(Counter)
        manufacturer_examples = defaultdict(lambda: defaultdict(list))

        mfr_ws = workbook["Manufacturer MPN Patterns"]
        for row in mfr_ws.iter_rows(min_row=2, values_only=True):
            manufacturer, mpn, exact, grouped, length, cls, _unused = (list(row) + [None] * 7)[:7]
            manufacturer = clean(manufacturer)
            mpn = clean(mpn)
            if not manufacturer or not mpn:
                continue
            grouped = clean(grouped) or grouped_mpn_pattern(mpn)
            manufacturer_rows += 1
            manufacturer_grouped_counts[manufacturer][grouped] += 1
            if len(manufacturer_examples[manufacturer][grouped]) < 3 and mpn not in manufacturer_examples[manufacturer][grouped]:
                manufacturer_examples[manufacturer][grouped].append(mpn)

        for manufacturer, counter in manufacturer_grouped_counts.items():
            top_patterns = [
                {
                    "pattern": pattern,
                    "count": count,
                    "examples": manufacturer_examples[manufacturer].get(pattern, []),
                }
                for pattern, count in counter.most_common(100)
            ]
            manufacturer_index[manufacturer] = {
                "mpn_count": sum(counter.values()),
                "top_grouped_patterns": top_patterns,
            }

        pattern_payload = {
            "version": 1,
            "source": input_path.name,
            "unique_mpn_rows": total_unique_rows,
            "pattern_rule": "Letters become A, digits become 0, and spaces/symbols are retained. Grouped patterns compress A/0 runs.",
            "exact_patterns": dict(exact_counts),
            "grouped_patterns": dict(grouped_counts),
            "character_classes": dict(class_counts),
            "lengths": dict(length_counts),
            "manufacturer_coverage": dict(manufacturer_coverage),
            "grouped_examples": dict(grouped_examples),
            "character_class_examples": dict(class_examples),
        }

        manufacturer_payload = {
            "version": 1,
            "source": input_path.name,
            "manufacturer_rows": manufacturer_rows,
            "manufacturer_count": len(manufacturer_index),
            "manufacturers": manufacturer_index,
        }

        pattern_path = output_dir / "mpn_patterns.v1.json.gz"
        manufacturer_path = output_dir / "manufacturer_pattern_index.v1.json.gz"
        write_json_gz(pattern_path, pattern_payload)
        write_json_gz(manufacturer_path, manufacturer_payload)

        self.stdout.write(self.style.SUCCESS(
            f"Built {total_unique_rows} unique MPN pattern rows and "
            f"{manufacturer_rows} manufacturer-pattern rows."
        ))
        self.stdout.write(str(pattern_path))
        self.stdout.write(str(manufacturer_path))
