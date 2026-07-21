import argparse
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd


MPN_START_RE = re.compile(r"(?=(?:^|\s)(?:AGILE|\d{5})[- ])")

KNOWN_MANUFACTURERS = [
    "3M (MINNESOTA MINING MFG)",
    "3L ELECTRONIC CORP.",
    "ABRACON CORPORATION",
    "ALLEGRO MICROSYSTEMS",
    "AMPHENOL FCI",
    "ANALOG DEVICES, INC.",
    "ANALOG DEVICES",
    "AVX/KYOCERA",
    "BOSSARD",
    "BOURNS",
    "BRADY WH CO",
    "BROADCOM CORPORATION",
    "COILCRAFT",
    "CTS CORP",
    "DIALIGHT CORP",
    "DIODES INC.",
    "DIODES",
    "ECLIPTEK",
    "EXXELIA",
    "FASTENAL",
    "FOX ELECTRONICS",
    "GENERAL ELECTRIC",
    "HARTING",
    "HERAEUS",
    "INFINEON TECHNOLOGIES AG",
    "INFINEON",
    "IRC",
    "KEMET",
    "KINGBRIGHT",
    "KOA",
    "KYOCERA AVX",
    "LITE-ON",
    "LITTELFUSE",
    "LOCTITE CORPORATION",
    "LOCTITE",
    "LUMEX",
    "MAXIM INTEGRATED PRODUCTS",
    "MCC (MICRO COMMERCIAL COMPONEN",
    "MICREL INC",
    "MICROCHIP TECHNO. (OLD TELCOM)",
    "MICROCHIP TECHNOLOGY INC",
    "MOLEX",
    "MURATA",
    "MYRRA",
    "NATIONAL SEMICONDUCTOR",
    "NEXPERIA B.V.",
    "NEXPERIA",
    "NIC COMPONENTS CORPORATION",
    "NIC COMPONENTS",
    "NXP",
    "OHMITE MANUFACTURING CO.",
    "ON SEMICONDUCTOR",
    "PANASONIC",
    "PERICOM, FORMERLY SARONIX",
    "PHOENIX CONTACT",
    "RALTRON (FORMERLY SHOWA)",
    "RCD COMPONENTS INC.",
    "RCD COMPONENTS INC",
    "RENESAS TECHNOLOGY CORP",
    "ROHM",
    "SAM'SUNG ELECTRO-MECHANICS CO",
    "SAMSUNG ELECTRO-MECHANICS CO",
    "SAMTEC INC",
    "SEI",
    "SEMIKRON INTERNATIONAL",
    "SHENNAN CIRCUITS",
    "SILICON STANDARD CORP.",
    "ST MICROELECTRONICS",
    "STANLEY ELECTRIC CO LTD",
    "SYE",
    "TAIWAN SEMICONDUCTOR",
    "TDK",
    "TE CONNECTIVITY",
    "TEXAS INSTRUMENTS",
    "TR FORMAC DIV OF TR FASTENINGS",
    "TT ELECTRONICS",
    "VALPEY FISHER DIV.OF VALTEC CO",
    "VCC LITE",
    "VISHAY",
    "YAGEO",
    "AVX",
]

KNOWN_MANUFACTURERS = sorted(set(KNOWN_MANUFACTURERS), key=len, reverse=True)


def clean(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def split_mpns(value):
    text = clean(value)
    if not text:
        return []

    starts = []
    for match in MPN_START_RE.finditer(text):
        start = match.start()
        if start < len(text) and text[start].isspace():
            start += 1
        starts.append(start)

    if not starts or starts[0] != 0:
        starts.insert(0, 0)

    starts = sorted(set(starts))
    parts = []
    for index, start in enumerate(starts):
        end = starts[index + 1] if index + 1 < len(starts) else len(text)
        part = text[start:end].strip()
        if part:
            parts.append(part)
    return parts


def split_manufacturers(value):
    remaining = clean(value)
    manufacturers = []

    while remaining:
        normalized = " ".join(remaining.split())
        matched = None
        for manufacturer in KNOWN_MANUFACTURERS:
            if normalized == manufacturer or normalized.startswith(manufacturer + " "):
                matched = manufacturer
                break

        if matched:
            manufacturers.append(matched)
            remaining = normalized[len(matched):].strip()
        else:
            # Keep moving rather than failing the whole row. The report flags
            # any rows where this fallback leads to count mismatches.
            token, _, rest = normalized.partition(" ")
            manufacturers.append(token)
            remaining = rest.strip()

    return manufacturers


def get_col_indexes(ws, header_row=4):
    indexes = {}
    for col in range(1, ws.max_column + 1):
        header = ws.cell(header_row, col).value
        if not header:
            continue
        indexes.setdefault(str(header).strip(), []).append(col)
    return indexes


def set_cell(ws, row, indexes, header, value, occurrence=0):
    cols = indexes.get(header, [])
    if occurrence < len(cols):
        ws.cell(row, cols[occurrence]).value = value


def map_bom_to_default(bom_path, template_path, output_path, report_path):
    bom_df = pd.read_excel(bom_path, sheet_name=0, header=1)
    bom_df = bom_df.dropna(how="all")

    workbook = openpyxl.load_workbook(template_path)
    worksheet = workbook.active
    indexes = get_col_indexes(worksheet, header_row=4)

    if worksheet.max_row >= 5:
        worksheet.delete_rows(5, worksheet.max_row - 4)

    output_row = 5
    report = {
        "source_rows": int(len(bom_df)),
        "output_rows": 0,
        "mismatch_rows": [],
    }

    for source_index, source in bom_df.iterrows():
        mpns = split_mpns(source.get("Manufacturer Equivalent part"))
        manufacturers = split_manufacturers(source.get("Manufacturer"))

        if not mpns:
            mpns = [""]

        if len(manufacturers) != len(mpns):
            report["mismatch_rows"].append({
                "bom_row": int(source_index) + 3,
                "mpn_count": len(mpns),
                "manufacturer_count": len(manufacturers),
                "manufacturer_cell": clean(source.get("Manufacturer")),
                "mpns": mpns,
                "parsed_manufacturers": manufacturers,
            })

        for mpn_index, mpn in enumerate(mpns):
            manufacturer = manufacturers[mpn_index] if mpn_index < len(manufacturers) else clean(source.get("Manufacturer"))

            set_cell(worksheet, output_row, indexes, "Item code", mpn)
            set_cell(worksheet, output_row, indexes, "MPN Code", mpn)
            set_cell(worksheet, output_row, indexes, "CPN Code", clean(source.get("Name")))
            set_cell(worksheet, output_row, indexes, "Item name", clean(source.get("Description")))
            set_cell(worksheet, output_row, indexes, "Description", clean(source.get("Description")))
            set_cell(worksheet, output_row, indexes, "Item type", "Raw material")
            set_cell(worksheet, output_row, indexes, "Measurement unit", clean(source.get("UOM")))

            set_cell(worksheet, output_row, indexes, "Specification name", "Reference Designator", occurrence=0)
            set_cell(worksheet, output_row, indexes, "Specification value", clean(source.get("Reference Designator")), occurrence=0)

            set_cell(worksheet, output_row, indexes, "Specification name", "Manufacturer", occurrence=1)
            set_cell(worksheet, output_row, indexes, "Specification value", manufacturer, occurrence=1)

            set_cell(worksheet, output_row, indexes, "Tag", manufacturer)

            output_row += 1

    report["output_rows"] = output_row - 5

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")

    return report


def main():
    parser = argparse.ArgumentParser(description="Map BOM 5.xlsx into Default Item template format.")
    parser.add_argument("--bom", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    report = map_bom_to_default(args.bom, args.template, args.output, args.report)
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
